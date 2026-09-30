"""Evidence-backed Russian SOC reports, shared by every output format."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import html
import json


def normalize_text(value):
    if isinstance(value, str):
        return value.replace("ё", "е").replace("Ё", "Е")
    if isinstance(value, list):
        return [normalize_text(item) for item in value]
    if isinstance(value, tuple):
        return tuple(normalize_text(item) for item in value)
    if isinstance(value, dict):
        return {normalize_text(key): normalize_text(item) for key, item in value.items()}
    return value


def number(value):
    if value is None:
        return "Нет данных"
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, int):
        return f"{value:,}".replace(",", " ")
    return str(value)


def percent(value, total):
    return f"{100 * value / total:.2f}%" if total else "Не рассчитано"


def evidence_digest(metrics):
    # Exclude generated presentation metadata, retaining all input evidence.
    data = {k: v for k, v in metrics.items() if k not in {"report_plan", "report_document"}}
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def build_document(metrics, title=None):
    total = metrics.get("total_events", 0)
    rules = metrics.get("rules") or {}
    integrity = metrics.get("integrity") or {}
    ingest = metrics.get("ingest") or {}
    meta = metrics.get("report_metadata") or {}
    context = metrics.get("source_context") or metrics.get("preflight") or {}
    findings = rules.get("findings") or []
    count = rules.get("findings_count", 0)
    sections = []

    def section(key, heading, paragraphs, headers=None, rows=None, chart=None, evidence=None):
        item = dict(id=key, title=heading, paragraphs=paragraphs, headers=headers or [],
                    rows=rows or [], chart=chart or [], evidence=evidence or [])
        sections.append(item)
        return item

    failures = []
    if integrity.get("ok") is not True:
        failures.append("Целостность входных данных не подтверждена.")
    for key in ("verified_rows", "expected_rows"):
        if integrity.get(key) != total:
            failures.append(f"Счетчик integrity.{key} отсутствует или не совпадает с total_events.")
    if rules.get("events_scanned") != total:
        failures.append("Покрытие проверки правилами не подтверждено для всех событий.")
    if sum((rules.get("findings_by_severity") or {}).values()) != count:
        failures.append("Распределение срабатываний по важности не сходится с общим числом.")
    if total <= 0:
        failures.append("Отсутствуют обработанные события.")

    cards = [
        {"id": "coverage", "text": f"Обработано {number(total)} событий. Проверено по манифесту: {number(integrity.get('verified_rows'))}.",
         "evidence": ["total_events", "integrity.verified_rows"], "action": "Сверить границы выборки, счетчики Splunk и манифест перед приемкой отчета."},
        {"id": "detections", "text": f"Зарегистрировано {number(count)} срабатываний правил. Одно событие может соответствовать нескольким правилам; количество подтвержденных инцидентов не установлено.",
         "evidence": ["rules.findings_count"], "action": "SOC L2: проверить контекст срабатываний высокого уровня и оформить результаты расследований."},
    ]
    if rules.get("findings_truncated"):
        cards.append({"id": "sample", "text": f"Сохранено {number(len(findings))} подробных записей из {number(count)} срабатываний ({percent(len(findings), count)}). Это ограниченная выборка, а не полный реестр.",
                      "evidence": ["rules.findings_truncated", "rules.findings_count"], "action": "Расширить реестр срабатываний с сохранением ссылок на исходные события; не экстраполировать состав выборки на весь массив."})
    requested = meta.get("target_events")
    if isinstance(requested, int) and requested != total:
        cards.append({"id": "reconciliation", "text": f"Целевой объем теста: {number(requested)}; фактическая выборка: {number(total)}; разница: {number(total - requested)}. Причина расхождения не установлена.",
                      "evidence": ["total_events", "report_metadata.target_events"], "action": "Владелец Splunk: проверить повторные отправки и границы тестового source; сохранить все записи до сверки."})
    plan = metrics.get("report_plan") or {}
    order = plan.get("selected_ids", []) if plan.get("status") == "validated" else []
    cards.sort(key=lambda c: order.index(c["id"]) if c["id"] in order else len(order))

    section("summary", "Резюме для руководства", [c["text"] for c in cards] + [
        "Результаты описывают наблюдаемую телеметрию и срабатывания детекторов. Подтверждение компрометации, ущерба и выполненных мер реагирования требует материалов расследования.",
        "Решение о приемке: отчет подготовлен автоматически и ожидает проверки ответственным аналитиком. Проверка чисел не заменяет подтверждение инцидентов."],
        ["Показатель", "Значение"], [
            ["Обработанные события", number(total)], ["Проверенные строки", number(integrity.get("verified_rows"))],
            ["Части экспорта", number(ingest.get("planned_chunks"))], ["Ошибки частей", number(ingest.get("failed_chunks"))],
            ["Срабатывания правил", number(count)], ["Подробности в выборке", number(len(findings))],
            ["Подтвержденные инциденты", "Не определены"], ["Проверка согласованности", "Пройдена" if not failures else "Есть замечания"]],
        evidence=["total_events", "rules", "integrity", "ingest"])
    section("scope", "Основание, объем и методика анализа", [
        "Цель отчета: представить результаты обработки журналов, наблюдаемые сигналы безопасности, качество доказательств и план дальнейшей проверки.",
        "Методика: экспорт по частям, проверка сохраненных строк, агрегирование полей, применение правил и подготовка аналитических разделов. Доли рассчитываются программно из исходных счетчиков.",
        "Тип данных: " + meta.get("dataset_description", "Не указан в метаданных запуска; производственное происхождение данных не подтверждено."),
        "Термины: событие - строка анализируемой выборки; срабатывание - совпадение с правилом; инцидент - подтвержденный результат расследования. Эти показатели не взаимозаменяемы.",
        "Отдельные распределения не доказывают связи между пользователем, узлом и типом события. Для корреляции нужны совместные поля и временные окна."],
        ["Параметр", "Значение"], [["Начало (включительно)", context.get("start", "Не задано")],
        ["Конец (исключительно)", context.get("end", "Не задано")], ["Индекс", context.get("index", "Не задан")],
        ["Источник выборки", context.get("source", "Не задан")], ["Заказчик", meta.get("customer", "Не указан")],
        ["Договор / основание", meta.get("contract", "Не указано")]], evidence=["source_context", "report_metadata"])

    timeline = metrics.get("timeline") or {}
    for key, heading in [("daily", "Динамика событий по дням"), ("hourly", "Распределение по часам суток")]:
        series = timeline.get(key) or []
        if series:
            if sum(v for _, v in series) != total:
                failures.append(f"Временной ряд {key} не согласован с общим числом событий.")
            peak = max(series, key=lambda row: row[1])
            section(key, heading, [
                "Часовой пояс: UTC. " + ("Показано число событий за каждые календарные сутки." if key == "daily" else "Суммированы одинаковые часы всех дней периода; это распределение по часу суток, а не последовательность последних 24 часов."),
                f"Максимум в представленном ряду: {peak[0]}, {number(peak[1])} событий ({percent(peak[1], total)} от всей выборки). Без независимого базового периода это не считается аномалией.",
                f"Метод: счетчики непересекающихся окон экспорта; для {number(timeline.get('boundary_chunks_read'))} частей, пересекающих границы часов, прочитаны временные метки всех строк и проверены хеши. Для остальных частей используется временной контракт исходного экспорта.",
                "Действие: сопоставить пики с расписанием систем и изменениями нагрузки. Падение объема требует проверки доступности источников, но само по себе не доказывает потерю журналов."],
                ["День UTC" if key == "daily" else "Час UTC", "События", "Доля от всех"],
                [[k, number(v), percent(v, total)] for k, v in series], series if key == "daily" else [], [f"timeline.{key}", "timeline.method"])

    dimensions = [
        ("sources", "Источники журналов", "top_sources", "Состав источников определяет, какие системы представлены в анализе. Имена файлов не подтверждают наличие всех необходимых классов событий.", "Сверить перечень источников с инвентаризацией и обязательными журналами."),
        ("types", "Типы телеметрии", "top_sourcetypes", "Тип источника отражает классификацию телеметрии. Наличие типа VPN или DNS само по себе не означает обнаружение угроз в соответствующей подсистеме.", "Проверить схемы нормализации и извлечение полей для каждого типа."),
        ("hosts", "Активность узлов инфраструктуры", "top_hosts", "Объем событий показывает вклад узла в выборку. Без истории и роли актива нельзя считать частоту аномальной или делать вывод о компрометации.", "Сопоставить узлы с реестром активов и проверить полноту поступления журналов."),
        ("users", "Учетные записи и контекст доступа", "top_users", "Агрегат по пользователям не разделяет успешные и неуспешные входы. Название учетной записи не доказывает фактические привилегии.", "Для приоритетных учетных записей проверить результат входа, целевой сервис и согласованные рабочие окна."),
        ("ips", "Сетевые источники", "top_src_ips", "IP-адреса в отчете маскируются. Одинаковые маски могут объединять множество адресов; число строк здесь не является числом уникальных устройств.", "В защищенном контуре сопоставить исходные адреса со срабатываниями и контекстом доступа."),
    ]
    for key, heading, field, meaning, action in dimensions:
        raw = metrics.get(field) or []
        combined = Counter()
        for label, value in raw:
            combined[str(label)] += value
        items = combined.most_common(50)
        subtotal = sum(v for _, v in items)
        if subtotal > total:
            failures.append(f"Сумма {field} превышает общий объем событий.")
        paragraphs = [meaning]
        if items:
            paragraphs.append(f"В таблице {number(len(items))} значений; суммарно {number(subtotal)} событий ({percent(subtotal, total)} от обработанной выборки). Лидер: {items[0][0]}, {number(items[0][1])} событий ({percent(items[0][1], total)}).")
            paragraphs.append("Остаток вне таблицы может включать пропуски поля и значения за пределами top-N. Его состав не установлен." if subtotal < total else "Сумма представленных значений совпадает с общим числом обработанных событий.")
        else:
            paragraphs.append("Распределение отсутствует во входных результатах. Система не восстанавливает значения по предположениям.")
        paragraphs.append("Действие: " + action)
        section(key, heading, paragraphs, ["Значение", "События", "Доля от всех"],
                [[k, number(v), percent(v, total)] for k, v in items], items[:12], [field, "total_events"])

    keywords = metrics.get("keyword_counts") or []
    section("signals", "Текстовые сигналы безопасности", [
        "Счетчик показывает наличие подстроки в тексте события. Это индикатор для поиска контекста, а не классификатор подтвержденных атак.",
        "Сигналы могут пересекаться: одно событие учитывается в нескольких строках. Складывать доли в общий процент инцидентов нельзя.",
        "Упоминание password не свидетельствует об утечке пароля. Упоминания malware и ransomware могут относиться к тестам, сообщениям защиты или описаниям, а не к заражению.",
        "Действие: проверить исходное сообщение, результат операции и правило детектирования; затем документировать решение аналитика."],
        ["Подстрока", "События", "Доля от всех"], [[k, number(v), percent(v, total)] for k, v in keywords],
        keywords, ["keyword_counts", "total_events"])
    severity = rules.get("findings_by_severity") or {}
    section("rules", "Результаты детектирования", [
        f"Загружено правил: {number(rules.get('rules_loaded'))}. Проверено событий: {number(rules.get('events_scanned'))}. Зарегистрировано срабатываний: {number(count)}.",
        "Важность задается конфигурацией правила. Она определяет порядок проверки сигнала, но не является оценкой подтвержденного ущерба.",
        f"Детальные записи сохранены для {number(len(findings))} срабатываний. Ограничение выборки: {number(rules.get('findings_truncated'))}.",
        "Полный счетчик срабатываний и ограниченная выборка имеют разные назначения. Распределение по идентификаторам ниже относится только к сохраненной выборке."],
        ["Уровень правила", "Срабатывания", "Доля срабатываний"],
        [[k, number(v), percent(v, count)] for k, v in severity.items()], list(severity.items()), ["rules"])

    by_rule = Counter(str(f.get("rule_id")) for f in findings)
    section("sample", "Реестр и примеры срабатываний", [
        "Ниже приведены примеры из сохраненной выборки, сгруппированные по правилу. Выборка ограничена порядком обработки и не является случайной или репрезентативной.",
        "Путь к части и номер строки позволяют найти событие в локальной выгрузке. Маскированные идентификаторы не обеспечивают однозначный поиск в Splunk.",
        "Статус всех примеров: требуется проверка. Действия блокировки, расследования или устранения по этим данным не подтверждены."],
        ["Правило", "В сохраненной выборке", "Доля выборки"],
        [[k, number(v), percent(v, len(findings))] for k, v in by_rule.most_common()], evidence=["rules.findings"])
    examples = []
    for rule_id in by_rule:
        examples.extend([f for f in findings if str(f.get("rule_id")) == rule_id][:2])
        if len(examples) >= 12:
            break
    for index, finding in enumerate(examples[:12], 1):
        ev = finding.get("evidence") or {}
        section(f"finding-{index}", f"Карточка {index}. {finding.get('rule_id', 'Без идентификатора')}", [
            "Наблюдаемый факт: сохраненная запись совпала с настроенным правилом. Текст исходного сообщения в отчет не включен.",
            "Проверка SOC: открыть указанную часть и строку, сопоставить событие с соседними событиями во времени, проверить разрешенность операции и зафиксировать вердикт.",
            "Вывод о реальном инциденте отсутствует. Повторяющиеся поля не доказывают дублирование: требуется сравнение исходных идентификаторов."],
            ["Поле доказательства", "Значение"], [["Правило", finding.get("rule_id")], ["Важность", finding.get("severity")]] +
            [[k, v] for k, v in ev.items()] + [["Часть", Path(finding.get("source_path", "")).name], ["Строка", number(finding.get("line"))]],
            evidence=["rules.findings"])

    section("coverage", "Полнота аналитики и ограничения", [
        "Отчет отражает доступные доказательства. Отсутствие метрики не означает отсутствие соответствующих событий или угроз.",
        "Динамика восстановлена по окнам экспорта с чтением пограничных частей." if timeline else "История по дням и часам не может быть восстановлена из общих частот. Временные метки отдельных срабатываний не описывают весь поток.",
        "Для следующего полного цикла нужны временные агрегаты, совместные распределения и результаты расследований. До их появления соответствующие выводы остаются недоступными."],
        ["Раздел анализа", "Доступность и предел вывода"], [
            ["Объем и распределения", "Доступны в пределах сохраненных агрегатов; см. предыдущие разделы."],
            ["Дни, часы, динамика", "Доступны, UTC; метод и границы описаны в соответствующих разделах." if timeline else "Нет полного временного ряда во входных метриках."],
            ["Успешные / неуспешные входы", "Есть текстовые сигналы и правила, нет полной нормализованной статистики результатов."],
            ["VPN, DNS, прокси", "Имена источников доступны; сессии, домены и решения политик не агрегированы."],
            ["Уязвимости, DLP, фишинг", "Нет специализированных результатов сканирования и расследований."],
            ["Уникальность событий", "Число строк подтверждается отдельно от дедупликации."],
            ["SLA и выполненные работы", "Нет журнала задач, времени реакции и подтверждений выполнения."],
            ["Детекторная точность", "Нет независимой размеченной контрольной выборки."]], evidence=["evaluation", "aggregation_limited"])
    models = metrics.get("model_comparison") or []
    section("models", "Модели и качество выводов", [
        "Таблица содержит фактически сохраненные результаты алгоритмов. Их шкалы различаются, поэтому сравнивать численные значения как единый рейтинг качества нельзя.",
        "rules_keywords: эвристический индекс по частотам подстрок. frequency_outliers: число редких узлов при заданном пороге. Ни один показатель не является процентом точности.",
        "Precision, recall, F1 и ROC-AUC в текущем запуске не измерены. Для измерения нужны разметка, разделение обучения и теста, единица оценки и зафиксированная версия детектора.",
        "Частотный профиль не является дообучением весов Qwen. Обновление профиля на тех же событиях не подтверждает способность распознавать новые атаки.",
        "Роль языковой модели: выбор порядка доказательных наблюдений. Модель получает компактные обезличенные карточки, возвращает только их идентификаторы; числа, формулировки фактов и все обязательные разделы формируются программно.",
        "Участие модели в этом документе: " + (f"{plan.get('model')}; ответ проверен, выбрано карточек: {len(order)}." if plan.get("status") == "validated" else "не подтверждено; использован детерминированный порядок."),
        "Такой режим ограничивает возможность выдуманных утверждений, но не проверяет истинность исходных логов и корректность самих правил."],
        ["Алгоритм", "Тип", "Сохраненный результат"],
        [[r.get("name"), r.get("type"), number(r.get("score"))] for r in models], evidence=["model_comparison", "evaluation", "report_plan"])
    section("integrity", "Контроль целостности и воспроизводимость", [
        "Проверка манифеста подтверждает согласованность сохраненных частей с зафиксированными хешами и счетчиками. Она не доказывает, что исходная система не потеряла события до экспорта.",
        "Совпадение количества строк не доказывает уникальность событий. Вопрос повторной отправки данных проверяется отдельно по устойчивому идентификатору.",
        "Замечания автоматического контроля: " + (" ".join(failures) if failures else "расхождений проверяемых счетчиков не обнаружено."),
        "Ограничение кардинальности агрегатов: " + number(metrics.get("aggregation_limited")) + "."],
        ["Проверка", "Результат"], [["Манифест", number(integrity.get("ok"))], ["Ожидалось строк", number(integrity.get("expected_rows"))],
        ["Проверено строк", number(integrity.get("verified_rows"))], ["Части экспорта", number(ingest.get("planned_chunks"))],
        ["Ошибки целостности", number(len(integrity.get("failures") or []))]], evidence=["integrity", "ingest"])
    section("actions", "План проверки и улучшений", [
        "Ниже перечислены предлагаемые действия. В отчете не утверждается, что они выполнены. Сроки и ответственные лица согласуются при приемке.",
        "Приоритет означает очередность проверки данных и сигналов, а не подтвержденную критичность инцидента."],
        ["Приоритет / роль", "Действие", "Критерий завершения"],
        [["1 / SOC и владелец SIEM", c["action"], "Результат сверки или расследования сохранен с доказательством."] for c in cards] + [
        ["2 / Инженер данных", "Добавить временные и совместные агрегаты в потоковый анализ.", "Суммы согласованы с обработанной выборкой; пропуски полей учтены."],
        ["2 / Владелец детекторов", "Подготовить размеченный тест, отдельно от обучения.", "Опубликованы матрица ошибок и метрики по классам."],
        ["3 / Руководитель SOC", "Проверить выводы, реквизиты и допуск к распространению.", "Зафиксированы рецензент, дата и решение о приемке."]], evidence=["rules", "integrity", "evaluation"])
    section("evidence", "Приложение. Доказательства и приемка", [
        "Источник чисел: приложенный JSON метрик. Пути evidence в разделах указывают поля этого файла. Контрольная сумма ниже вычислена по каноническому JSON входных доказательств без служебного плана отчета.",
        "SHA-256 доказательств: " + evidence_digest(metrics),
        "Состав комплекта: PDF для чтения, HTML и Markdown с тем же содержанием, JSON метрик и структурированный документ с результатами автоматической проверки.",
        "Проверка аналитиком: ФИО, дата и решение не заполнены. Документ не содержит электронной подписи и не подтверждает выполнение договорных обязательств."],
        ["Артефакт", "Назначение"], [["Метрики", "Исходные счетчики, правила, выборка срабатываний."],
        ["Манифест", str(integrity.get("manifest", "Не указан"))],
        ["Версия структуры", "soc-report-v2"], ["Статус", "На проверке аналитика"]], evidence=["integrity.manifest"])
    return normalize_text(dict(schema="soc-report-v2", title="Отчет по результатам анализа событий информационной безопасности",
                customer=meta.get("customer", "Тестовый контур SOC"), period=f"{context.get('start', 'Не задано')} - {context.get('end', 'Не задано')}",
                generated=datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M UTC"), total=total,
                sections=sections, cards=cards, status="На проверке аналитика", failures=failures,
                evidence_sha256=evidence_digest(metrics)))


def markdown_document(doc):
    def cell(value):
        return html.escape(str(normalize_text(value))).replace("|", "&#124;").replace("\n", " ").replace("\r", " ")
    lines = [f"# {doc['title']}", "", cell(doc["customer"]), cell(doc["period"]), doc["status"], ""]
    for i, section in enumerate(doc["sections"], 1):
        lines.extend([f"## {i}. {section['title']}", ""])
        lines.extend(cell(p) + "\n" for p in section["paragraphs"])
        if section["rows"]:
            lines += ["| " + " | ".join(map(cell, section["headers"])) + " |", "| " + " | ".join("---" for _ in section["headers"]) + " |"]
            lines += ["| " + " | ".join(map(cell, row)) + " |" for row in section["rows"]]
        lines.extend(["", "Evidence: " + ", ".join(section["evidence"]), ""])
    return "\n".join(lines)


def html_document(doc):
    esc = lambda v: html.escape(str(normalize_text(v)))
    body = f"<h1>{esc(doc['title'])}</h1><p>{esc(doc['customer'])}</p><p>{esc(doc['period'])}</p><p>{esc(doc['status'])}</p>"
    for i, s in enumerate(doc["sections"], 1):
        body += f"<section><h2>{i}. {esc(s['title'])}</h2>" + "".join(f"<p>{esc(p)}</p>" for p in s["paragraphs"])
        if s["rows"]:
            body += "<table><thead><tr>" + "".join(f"<th>{esc(c)}</th>" for c in s["headers"]) + "</tr></thead><tbody>"
            body += "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in s["rows"]) + "</tbody></table>"
        body += f"<p class='evidence'>Evidence: {esc(', '.join(s['evidence']))}</p></section>"
    return "<!doctype html><html lang='ru'><meta charset='utf-8'><title>SOC report</title><style>body{max-width:980px;margin:40px auto;padding:0 24px;font:16px/1.6 Georgia,serif;color:#20252c}h1{font-size:28px}h2{font-size:23px}section{margin-top:40px}table{width:100%;border-collapse:collapse;font:14px/1.5 sans-serif}th,td{padding:10px;border:1px solid #bbb;text-align:left;overflow-wrap:anywhere}th{background:#edf1f2}.evidence{font:12px monospace;color:#555}@media print{section{break-before:page}thead{display:table-header-group}}</style>" + body + "</html>"


def pdf_document(doc, output):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, PageBreak, LongTable, TableStyle
    from reportlab.platypus.tableofcontents import TableOfContents
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.graphics.shapes import Drawing, Group, Line, Polygon, Rect, String
    fonts = Path(__file__).parent / "fonts"
    def register_font(name, candidates):
        if name in pdfmetrics.getRegisteredFontNames():
            return name
        for candidate in candidates:
            path = Path(candidate)
            if path.exists():
                pdfmetrics.registerFont(TTFont(name, str(path)))
                return name
        return "Helvetica-Bold" if name.endswith("Bold") else "Helvetica"

    regular = register_font("SOCSerif", [
        fonts / "DejaVuSerif.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"),
        Path("C:/Windows/Fonts/DejaVuSerif.ttf"),
        fonts / "DejaVuSans.ttf",
    ])
    bold = register_font("SOCSerif-Bold", [
        fonts / "DejaVuSerif-Bold.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"),
        Path("C:/Windows/Fonts/DejaVuSerif-Bold.ttf"),
        fonts / "DejaVuSans-Bold.ttf",
    ])
    sans = register_font("SOCSans", [fonts / "DejaVuSans.ttf", Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")])
    sans_bold = register_font("SOCSans-Bold", [fonts / "DejaVuSans-Bold.ttf", Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")])
    pdfmetrics.registerFontFamily("SOCSerif", normal=regular, bold=bold, italic=regular, boldItalic=bold)
    ink, muted, accent = colors.HexColor("#202020"), colors.HexColor("#4f4f4f"), colors.HexColor("#6f6f6f")
    grid = colors.HexColor("#202020")
    light = colors.HexColor("#f2f2f2")
    styles = {
        "body": ParagraphStyle("Body", fontName=regular, fontSize=11, leading=17, firstLineIndent=25, alignment=4, spaceAfter=8, textColor=ink, splitLongWords=True),
        "heading": ParagraphStyle("Section", fontName=bold, fontSize=14, leading=18, spaceAfter=13, keepWithNext=True, textColor=ink),
        "cover_heading": ParagraphStyle("CoverSubtitle", fontName=bold, fontSize=14, leading=20, alignment=1, spaceAfter=12, textColor=ink),
        "small": ParagraphStyle("Small", fontName=regular, fontSize=8.5, leading=12, spaceAfter=8, textColor=muted),
        "caption": ParagraphStyle("Caption", fontName=regular, fontSize=10, leading=13, leftIndent=10, spaceAfter=6, textColor=ink),
        "cell": ParagraphStyle("Cell", fontName=regular, fontSize=9.2, leading=12.5, alignment=1, textColor=ink, splitLongWords=True),
        "header": ParagraphStyle("Header", fontName=bold, fontSize=9.2, leading=12.5, alignment=1, textColor=ink),
        "title": ParagraphStyle("Cover", fontName=bold, fontSize=18, leading=25, alignment=1, spaceAfter=16, textColor=ink),
    }
    def p(value, style="body"):
        return Paragraph(html.escape(str(normalize_text(value))), styles[style])

    def shorten(text, font_name, font_size, max_width):
        value = " ".join(str(normalize_text(text)).split())
        while value and pdfmetrics.stringWidth(value, font_name, font_size) > max_width:
            value = value[:-2].rstrip()
        return value + ("." if value and not str(text).startswith(value) else "")

    class Document(BaseDocTemplate):
        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and flowable.style.name == "Section":
                self.notify("TOCEntry", (0, flowable.getPlainText(), self.page))

    class NumberedCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self._total_pages = total_pages
                super().showPage()
            super().save()

    def page_chrome(canvas, document):
        canvas.saveState()
        left, top, header_h = 50, 785, 58
        logo_w, page_w = 104, 78
        middle_w = 495 - logo_w - page_w
        canvas.setLineWidth(0.7)
        canvas.setStrokeColor(grid)
        canvas.rect(left, top - header_h, 495, header_h, stroke=1, fill=0)
        canvas.line(left + logo_w, top - header_h, left + logo_w, top)
        canvas.line(left + logo_w + middle_w, top - header_h, left + logo_w + middle_w, top)
        canvas.line(left + logo_w + middle_w, top - 20, left + 495, top - 20)

        canvas.setFont(sans_bold, 13)
        canvas.setFillColor(colors.HexColor("#26314a"))
        canvas.drawCentredString(left + logo_w / 2, top - 27, "SOC AI")
        canvas.setFont(sans, 5.8)
        canvas.setFillColor(muted)
        canvas.drawCentredString(left + logo_w / 2, top - 39, "security operations report")

        canvas.setFont(regular, 9.5)
        canvas.setFillColor(ink)
        customer = shorten(doc["customer"], regular, 9.5, middle_w - 24)
        period = shorten("Период: " + doc["period"], regular, 8.5, middle_w - 24)
        canvas.drawCentredString(left + logo_w + middle_w / 2, top - 26, customer)
        canvas.setFont(regular, 8.5)
        canvas.drawCentredString(left + logo_w + middle_w / 2, top - 41, period)

        canvas.setFont(regular, 9.5)
        canvas.drawCentredString(left + logo_w + middle_w + page_w / 2, top - 14, "страница")
        canvas.setFont(regular, 9.5)
        canvas.drawCentredString(left + logo_w + middle_w + page_w / 2, top - 39, str(document.page))
        canvas.setStrokeColor(colors.HexColor("#b7b7b7"))
        canvas.line(50, 42, 545, 42)
        canvas.setFont(sans, 7.2)
        canvas.setFillColor(muted)
        canvas.drawString(50, 29, "Для ограниченного распространения. Автоматический отчет подлежит проверке аналитиком.")
        canvas.drawRightString(545, 29, "soc-report-v2")
        canvas.restoreState()

    document = Document(str(output), pagesize=(595.28, 841.89), title=doc["title"], author="SOC AI Platform")
    document.addPageTemplates(PageTemplate(id="normal", frames=[Frame(50, 60, 495, 650, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)], onPage=page_chrome))
    story = [Spacer(1, 185), p("ОТЧЕТ", "title"), p(doc["title"], "cover_heading"), Spacer(1, 118),
             p("Заказчик: " + doc["customer"], "body"), p("Отчетный период: " + doc["period"], "body"),
             p("Проанализировано событий: " + number(doc["total"]), "body"), Spacer(1, 80),
             p("Сформирован: " + doc["generated"], "small"), PageBreak(), p("Содержание", "cover_heading")]
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("Contents", fontName=regular, fontSize=9.5, leading=13, spaceBefore=4)]
    story += [toc]

    def table_widths(section, columns):
        if columns == 2:
            return [247.5, 247.5]
        if section["id"] == "actions":
            return [86, 244, 165]
        return [225, 135, 135]

    def three_d_chart(section, chart):
        chart = chart[:6 if len(section["rows"]) > 10 else 10]
        if not chart or max(v for _, v in chart) <= 0:
            return []
        maximum = max(v for _, v in chart)
        row_h = 24
        height = len(chart) * row_h + 16
        drawing = Drawing(495, height)
        label_w, bar_x, bar_w, depth = 151, 164, 230, 6
        top_color = colors.HexColor("#9c9c9c")
        side_color = colors.HexColor("#585858")
        front_color = colors.HexColor("#767676")
        for j, (label, value) in enumerate(chart):
            y = height - row_h * (j + 1) + 3
            short = shorten(label, regular, 8, label_w - 4)
            width = max(3, bar_w * value / maximum)
            drawing.add(String(0, y + 2, short, fontName=regular, fontSize=8, fillColor=ink))
            group = Group()
            group.add(Rect(bar_x, y, width, 11, fillColor=front_color, strokeColor=colors.HexColor("#404040"), strokeWidth=.25))
            group.add(Polygon([bar_x, y + 11, bar_x + depth, y + 16, bar_x + width + depth, y + 16, bar_x + width, y + 11],
                              fillColor=top_color, strokeColor=colors.HexColor("#404040"), strokeWidth=.25))
            group.add(Polygon([bar_x + width, y, bar_x + width + depth, y + 5, bar_x + width + depth, y + 16, bar_x + width, y + 11],
                              fillColor=side_color, strokeColor=colors.HexColor("#404040"), strokeWidth=.25))
            drawing.add(group)
            drawing.add(String(492, y + 2, number(value), textAnchor="end", fontName=regular, fontSize=8, fillColor=ink))
        drawing.add(Line(bar_x, 4, bar_x + bar_w + depth, 4, strokeColor=colors.HexColor("#707070"), strokeWidth=.35))
        return [p(f"Рисунок. {section['title']}; 3D-гистограмма, показано значений: {len(chart)}", "caption"), drawing, Spacer(1, 10)]

    for i, s in enumerate(doc["sections"], 1):
        story += [PageBreak(), p(f"{i}. {s['title']}", "heading")]
        story += [p(text) for text in s["paragraphs"]]
        if s["rows"]:
            story += [p(f"Таблица №{i}. {s['title']}", "caption")]
            n = len(s["headers"])
            widths = table_widths(s, n)
            table = LongTable([[p(c, "header") for c in s["headers"]]] + [[p(c, "cell") for c in row] for row in s["rows"]], colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), accent), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, light]),
                ("GRID", (0,0), (-1,-1), .55, grid), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
                ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
                ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5)]))
            story += [table, Spacer(1, 14)]
        story += three_d_chart(s, s["chart"])
        story += [p("Evidence: " + ", ".join(s["evidence"]), "small")]
    document.multiBuild(story, canvasmaker=NumberedCanvas)


def docx_document(doc, output):
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def set_cell_border(cell, color="606060", size="6"):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.first_child_found_in("w:tcBorders")
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), size)
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), color)

    def shade(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def add_page_number(paragraph):
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.extend([fld_begin, instr, fld_end])

    def paragraph(text="", style=None, align=None):
        item = document.add_paragraph(style=style)
        if align is not None:
            item.alignment = align
        item.add_run(str(normalize_text(text)))
        return item

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    for style_name in ("Normal", "Body Text"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    header_table = section.header.add_table(rows=1, cols=3, width=Cm(17.4))
    header_table.autofit = False
    widths = [Cm(3.6), Cm(10.7), Cm(3.1)]
    for index, width in enumerate(widths):
        header_table.columns[index].width = width
    cells = header_table.rows[0].cells
    cells[0].text = "SOC AI\nsecurity operations report"
    cells[1].text = f"{doc['customer']}\nПериод: {doc['period']}"
    cells[2].text = "страница\n"
    add_page_number(cells[2].paragraphs[-1])
    for cell in cells:
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.text = "Для ограниченного распространения. Автоматический отчет подлежит проверке аналитиком."
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.runs[0].font.size = Pt(8)

    document.add_paragraph()
    document.add_paragraph()
    paragraph("ОТЧЕТ", "Title", WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc["title"], "Heading 1", WD_ALIGN_PARAGRAPH.CENTER)
    document.add_paragraph()
    document.add_paragraph()
    paragraph("Заказчик: " + doc["customer"])
    paragraph("Отчетный период: " + doc["period"])
    paragraph("Проанализировано событий: " + number(doc["total"]))
    paragraph("Сформирован: " + doc["generated"])
    document.add_page_break()

    paragraph("Содержание", "Heading 1", WD_ALIGN_PARAGRAPH.CENTER)
    for index, section_data in enumerate(doc["sections"], 1):
        paragraph(f"{index}. {section_data['title']}")
    document.add_page_break()

    for index, section_data in enumerate(doc["sections"], 1):
        paragraph(f"{index}. {section_data['title']}", "Heading 1")
        for item in section_data["paragraphs"]:
            paragraph(item)
        if section_data["rows"]:
            paragraph(f"Таблица №{index}. {section_data['title']}")
            table = document.add_table(rows=1, cols=len(section_data["headers"]))
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, section_data["headers"]):
                cell.text = str(normalize_text(value))
                shade(cell, "707070")
                set_cell_border(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.bold = True
                        run.font.size = Pt(9)
            for row_index, row in enumerate(section_data["rows"], 1):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(normalize_text(value))
                    set_cell_border(cell)
                    if row_index % 2 == 0:
                        shade(cell, "F2F2F2")
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(9)
            document.add_paragraph()
        if section_data["chart"]:
            paragraph(f"Рисунок. {section_data['title']}; график доступен в PDF-версии отчета.")
        paragraph("Evidence: " + ", ".join(section_data["evidence"]))
        if index != len(doc["sections"]):
            document.add_section(WD_SECTION_START.NEW_PAGE)

    document.save(str(output))
